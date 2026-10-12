"""Build docs/PROJECT_REPORT.docx (and .pdf if Microsoft Word is installed) from PROJECT_REPORT.md.

    python scripts/build_report.py

Handles the Markdown subset used by the report: headings, paragraphs, bullet and
numbered lists, pipe tables, images, **bold**, *italic* and `code`.
"""

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
SRC = DOCS / "PROJECT_REPORT.md"
OUT = DOCS / "PROJECT_REPORT.docx"
PDF = DOCS / "PROJECT_REPORT.pdf"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("*"):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(set(c) <= set("-: ") for c in r)]
    table = doc.add_table(rows=len(cells), cols=len(cells[0]))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j, text in enumerate(row[: len(cells[0])]):
            par = table.cell(i, j).paragraphs[0]
            add_inline(par, text)
            for run in par.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def build():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if m:
            img = DOCS / m.group(2)
            if img.exists():
                doc.add_picture(str(img), width=Cm(16))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(m.group(1))
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(80, 80, 80)
        elif line.startswith("# "):
            t = doc.add_heading(line[2:], level=0)
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif re.match(r"^\s*- ", line):
            add_inline(doc.add_paragraph(style="List Bullet"), line.strip()[2:])
        elif re.match(r"^\d+\. ", line):
            add_inline(doc.add_paragraph(style="List Number"), line.split(". ", 1)[1])
        elif line.strip():
            # merge wrapped paragraph lines
            text = line.strip()
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||!\[|\s*- |\d+\. )", lines[i + 1]):
                i += 1
                text += " " + lines[i].strip()
            par = doc.add_paragraph()
            if text.startswith("x̄") or text.startswith("M(x") or text.startswith("p_t ="):
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(par, text)
        i += 1
    doc.save(OUT)
    print("wrote", OUT)


def export_pdf() -> bool:
    """Export with Word via COM (PowerShell). Returns False if Word is unavailable."""
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{OUT}'); $d.SaveAs([ref]'{PDF}', [ref]17); "
        "$d.Close(); $w.Quit()"
    )
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, timeout=180,
                       capture_output=True)
    except (OSError, subprocess.SubprocessError) as exc:
        print("PDF export skipped:", exc)
        return False
    ok = PDF.exists()
    print("wrote", PDF if ok else "no PDF produced")
    return ok


if __name__ == "__main__":
    build()
    if "--no-pdf" not in sys.argv:
        export_pdf()
